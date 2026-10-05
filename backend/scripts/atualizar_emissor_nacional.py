"""Atualiza app/data/emissor_nacional.json a partir da planilha oficial de
monitoramento das adesões à NFS-e (Receita Federal).

Uso:
    1. Baixe a planilha mais nova em
       https://www.gov.br/nfse/pt-br/municipios/monitoramento-adesoes
       (arquivo municipios-aderentes-AAAAMMDD.xlsx).
    2. cd backend && python3 scripts/atualizar_emissor_nacional.py caminho/da/planilha.xlsx

Entra na lista o município "Conveniado Ativo" com AderenteEmissorNacional =
Sim. A planilha não traz o código do IBGE: o casamento é por UF + nome.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.services.municipios import _normalizar, _tabela  # noqa: E402

# Nomes que a planilha escreve diferente da tabela do IBGE.
APELIDOS = {
    ("GO", "bom jesus"): "5203500", ("BA", "lagedo do tabocal"): "2919058", ("RN", "boa saude"): "2405306",
    ("SE", "amparo de sao francisco"): "2800100", ("TO", "sao valerio da natividade"): "1720499",
    ("MG", "sao thome das letras"): "3165206", ("PR", "munhoz de mello"): "4116307",
}


def main(caminho: str) -> None:
    lista, _ = _tabela()
    por_nome = {(m["uf"], m["_busca"]): m["codigo"] for m in lista}
    linhas = load_workbook(caminho, read_only=True).worksheets[0].iter_rows(values_only=True)
    cabecalho = [str(c or "").strip() for c in next(linhas)]
    col = {nome: cabecalho.index(nome) for nome in ("UF", "NomeMunicipio", "StatusConvenioSEFIN", "AderenteEmissorNacional")}
    codigos, sem_par = set(), []
    for linha in linhas:
        if str(linha[col["AderenteEmissorNacional"]] or "").strip() != "Sim" or "Ativo" not in str(linha[col["StatusConvenioSEFIN"]] or ""):
            continue
        chave = (str(linha[col["UF"]]).strip().upper(), _normalizar(str(linha[col["NomeMunicipio"]])))
        codigo = por_nome.get(chave) or APELIDOS.get(chave)
        (codigos.add(codigo) if codigo else sem_par.append(chave))
    data = re.search(r"(\d{4})(\d{2})(\d{2})", Path(caminho).name)
    destino = Path(__file__).resolve().parent.parent / "app" / "data" / "emissor_nacional.json"
    destino.write_text(json.dumps({
        "fonte": "https://www.gov.br/nfse/pt-br/municipios/monitoramento-adesoes",
        "atualizado_em": "-".join(data.groups()) if data else None, "codigos": sorted(codigos),
    }, separators=(",", ":")), encoding="utf-8")
    print(f"{len(codigos)} municípios usam o Emissor Nacional.")
    if sem_par:
        print("Sem par na tabela do IBGE (acrescente em APELIDOS):", sem_par)


if __name__ == "__main__":
    main(sys.argv[1])
